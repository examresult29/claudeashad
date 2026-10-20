#!/usr/bin/env python3
"""Build BD-e-Return-Calculator.xlsx.

Sheets:
  Input   - everything you type (peach cells, blue text); results and checks at the top
  Rules   - tax-free limits, slabs, minimum tax, rebate, surcharge per rules profile
  Calc    - every figure of the return as a live formula (named v_...)
  P01-P12 - the 12 printed pages, laid out like the sample return

Run:  python3 excel/build_workbook.py [output.xlsx]
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SKILL = ROOT / "bd-e-return" / "skills" / "e-return"
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(HERE))

from ereturn.calc import GOVT_ROWS, INVEST_ROWS, PRIVATE_ROWS, TAXFREE_ROWS  # noqa: E402
from ereturn.fields import PAGES  # noqa: E402
from layout import PADL, PADR, Item, build_grid  # noqa: E402

TEMPLATE = json.loads((SKILL / "ereturn" / "template.json").read_text(encoding="utf-8"))
PROFILES = ["ay2026-27", "sample-template", "fo2025"]
RULES = {p: json.loads((SKILL / "ereturn" / "rules" / f"{p}.json").read_text(encoding="utf-8")) for p in PROFILES}
EXAMPLE = json.loads((SKILL / "examples" / "sample-reproduction.json").read_text(encoding="utf-8"))
EMBLEM = SKILL / "ereturn" / "assets" / "emblem.png"

PEACH = PatternFill("solid", fgColor="FDE9D9")
GREEN = "006A4E"
BLUE = "0000FF"
MONEY = '#,##0;(#,##0);"-"'
PAGE_NAMES = ["P01 Acknowledgement", "P02 Return page 1", "P03 Return page 2", "P04 Return page 3",
              "P05 Schedule 1", "P06 Schedule 2-3", "P07 Schedule 4-5", "P08 IT-10BB",
              "P09 IT-10B page 1", "P10 IT-10B page 2", "P11 Tax calculation", "P12 Tax free income"]
# The sample printed these pages shrunk (fonts of 11 and 12 pt come out as 9.24 and 10.08 at
# 84%); the sheets are laid out at full size and print at the same scale, so fonts are exact.
PAGE_SCALE = {8: 0.84, 9: 0.84, 10: 0.90, 11: 0.90}
EXCEL_FONT = {"serif": ("Times New Roman", False), "serif-bold": ("Times New Roman", True),
              "calibri": ("Calibri", False), "calibri-bold": ("Calibri", True),
              "sans": ("Arial", False), "sans-bold": ("Arial", True),
              "taka": ("Tahoma", False), "symbol": ("MS Gothic", False)}

# Column widths are counted in digits of the Normal style's font. Excel rounds a digit to whole
# pixels and LibreOffice does not, so the two disagree for Calibri 11 (7 px vs 7.4 px) and pages
# print about 6% short in LibreOffice. Times New Roman 12 has digits of exactly 8 px in both.
NORMAL_FONT = Font(name="Times New Roman", size=12, family=1)
DIGIT_PX = 8
CHARS = 7 / DIGIT_PX  # the Input, Rules and Calc widths below were chosen in Calibri 11 digits

# Text widths (only used to size merged cells), with the metric twins of the Excel fonts.
for nm, path in (("Carlito", "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf"),
                 ("Carlito-Bold", "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf")):
    try:
        pdfmetrics.registerFont(TTFont(nm, path))
    except Exception:
        pass
RL_FONT = {"serif": "Times-Roman", "serif-bold": "Times-Bold", "sans": "Helvetica", "sans-bold": "Helvetica-Bold",
           "calibri": "Carlito", "calibri-bold": "Carlito-Bold"}


def text_width(s: str, role: str, size: float) -> float:
    font = RL_FONT.get(role)
    if not font:
        return len(s) * size * 0.6
    try:
        return pdfmetrics.stringWidth(s, font, size)
    except Exception:
        return pdfmetrics.stringWidth(s, "Helvetica", size) * 0.92


def excel_size(size: float) -> float:
    return round(size * 2) / 2


def print_wide(ws, landscape=False):
    """Working sheets print one page wide, as many pages tall as they need."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def add_name(wb, name: str, ref: str):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def ex(path: str, default=None):
    """Value from the example input, by dotted path."""
    cur = EXAMPLE
    for part in path.split("."):
        if isinstance(cur, list):
            part = int(part)
            if part >= len(cur):
                return default
            cur = cur[part]
        elif isinstance(cur, dict):
            if part not in cur:
                return default
            cur = cur[part]
        else:
            return default
    return cur


# =================================================================== Input sheet
YESNO = ["Yes", "No"]


class InputSheet:
    def __init__(self, wb):
        self.wb = wb
        self.ws = wb.active
        self.ws.title = "Input"
        self.row = 1
        self.dv = {}
        ws = self.ws
        for col, w in (("A", 2), ("B", 50), ("C", 22), ("D", 22), ("E", 22), ("F", 64)):
            ws.column_dimensions[col].width = w * CHARS
        ws.sheet_view.showGridLines = False

    def validation(self, options):
        key = ",".join(options)
        if key not in self.dv:
            dv = DataValidation(type="list", formula1='"' + key + '"', allow_blank=True)
            self.ws.add_data_validation(dv)
            self.dv[key] = dv
        return self.dv[key]

    def heading(self, text, note=None):
        self.row += 1
        c = self.ws.cell(self.row, 2, text)
        c.font = Font(name="Times New Roman", size=14, bold=True, color=GREEN)
        if note:
            self.note(note)
        self.row += 1

    def sub(self, text):
        c = self.ws.cell(self.row, 2, text)
        c.font = Font(name="Arial", size=10, bold=True)
        self.row += 1

    def cell(self, col, name, kind, value, options=None):
        c = self.ws.cell(self.row, col)
        c.fill = PEACH
        c.font = Font(name="Arial", size=10, color=BLUE)
        c.border = Border(bottom=Side(style="thin", color="C9A98F"))
        if kind == "num":
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")
        elif kind == "date":
            c.number_format = "dd/mm/yyyy"
            c.alignment = Alignment(horizontal="left")
        elif kind == "text":
            c.number_format = "@"
        if options:
            self.validation(options).add(c.coordinate)
        if kind == "num" and isinstance(value, str):
            value = float(value.replace(",", "")) if value.replace(",", "").strip() else None
        if value not in (None, ""):
            if kind == "date":
                from datetime import date
                y, m, d = (int(x) for x in str(value).split("-"))
                c.value = date(y, m, d)
            elif kind == "text":
                c.value = str(value)
            else:
                c.value = value
        add_name(self.wb, name, f"Input!${get_column_letter(col)}${self.row}")
        return c

    def field(self, label, name, kind="num", value=None, note=None, options=None):
        b = self.ws.cell(self.row, 2, label)
        b.font = Font(name="Arial", size=10)
        self.cell(3, name, kind, value, options)
        if note:
            self.note(note)
        self.row += 1

    def note(self, text):
        n = self.ws.cell(self.row, 6, text)
        n.font = Font(name="Arial", size=9, color="5D6B64")
        n.alignment = Alignment(wrap_text=True, vertical="top")  # rows grow to fit when opened


def build_input(wb):
    s = InputSheet(wb)
    ws = s.ws
    ws["B1"] = "BD e-Return calculator"
    ws["B1"].font = Font(name="Times New Roman", size=20, bold=True, color=GREEN)
    ws["B2"] = ("Fill in the peach cells (blue text). Everything else is calculated. "
                "The workbook opens with the sample return's figures as an example: replace them with your taxpayer's.")
    ws["B2"].font = Font(name="Arial", size=10)
    ws["B3"] = ("To print the return: click the tab \"P01 Acknowledgement\", hold Shift, click \"P12 Tax free income\", "
                "then press Ctrl+P (for a PDF: File > Save As > PDF, Options: \"Active sheet(s)\"). "
                "The 12 pages print in order, one A4 page each.")
    ws["B3"].font = Font(name="Arial", size=10, bold=True)
    for r in (2, 3):  # long lines wrap inside the sheet's width, so the sheet prints one page wide
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 27
    s.row = 5
    results = s.row

    # ---------------------------------------------------------------- results
    s.heading("Results", "Live from the Calc sheet.")
    for label, formula in (
        ("Total income (return line 11)", "=v_r11"), ("Tax-exempt income (line 26)", "=v_r26"),
        ("Tax-free limit used", "=v_threshold"), ("Tax on slabs (line 12)", "=v_r12"),
        ("Rebate on investment (line 13)", "=v_r13"), ("Minimum tax (line 15)", "=v_r15"),
        ("Tax payable (line 16)", "=v_r16"), ("Surcharge, delay interest, final tax (lines 17-18)", "=v_r17+v_r18"),
        ("Total amount payable (line 19)", "=v_r19"), ("Already paid and adjusted (lines 20-22)", "=v_r20+v_r21+v_r22"),
        ("Pay with the return (line 23)", "=v_r23"), ("Deficit or (excess) (line 25)", "=v_r25"),
        ("Net wealth (IT-10B line 5)", "=v_w_5"),
        ("Balancing figure", '=IF(v_mode="gift","Receipt of gift and others: ",IF(v_mode="cash","Cash and bank: ","Not balanced automatically: "))'
                             '&TEXT(IF(v_mode="gift",v_w_1c,v_a_k),"#,##0")'),
    ):
        ws.cell(s.row, 2, label).font = Font(name="Arial", size=10, bold=label.startswith(("Pay with", "Total amount")))
        c = ws.cell(s.row, 3, formula)
        c.number_format = MONEY
        c.font = Font(name="Arial", size=10, bold=label.startswith(("Pay with", "Total amount")))
        if label.startswith("Pay with"):
            c.fill = PatternFill("solid", fgColor="DCEFE7")
        s.row += 1
    s.heading("Checks", "Messages appear here when something needs attention.")
    checks = ("v_chk_name", "v_chk_tin", "v_chk_gift", "v_chk_cash", "v_chk_balance", "v_chk_biz",
              "v_chk_deficit", "v_chk_min")
    ok = ws.cell(s.row, 2, f'=IF({"&".join(checks)}="","No problems found in the figures.","")')
    ok.font = Font(name="Arial", size=10, color=GREEN)
    s.row += 1
    for f in checks + ("v_chk_rules",):
        c = ws.cell(s.row, 2, "=" + f)
        c.font = Font(name="Arial", size=10, color="B3261E")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=s.row, start_column=2, end_row=s.row, end_column=6)
        if f == "v_chk_rules":
            ws.row_dimensions[s.row].height = 27
        s.row += 1
    ws.freeze_panes = None
    print_wide(ws, landscape=True)

    # ---------------------------------------------------------------- settings
    s.heading("Return settings")
    s.field("Assessment year", "in_ay", "text", ex("assessment_year", "2026-2027"),
            "The income year ends on 30 June of the first year.")
    s.field("Tax rules", "in_profile", "text", ex("rules_profile", "ay2026-27"),
            "ay2026-27 = AY 2026-27 as reported for the Finance Act 2026; sample-template = the sample's older slabs; "
            "fo2025 = Finance Ordinance 2025. See the Rules sheet.", PROFILES)

    # ---------------------------------------------------------------- taxpayer
    s.heading("Taxpayer", "Printed on the return's first page and on every schedule.")
    tp = "taxpayer."
    s.field("Name of the assessee", "in_name", "text", ex(tp + "name"))
    s.field("TIN (12 digits)", "in_tin", "text", ex(tp + "tin"))
    s.field("NID or passport number", "in_nid", "text", ex(tp + "nid"))
    s.field("Circle", "in_circle", "text", ex(tp + "circle"))
    s.field("Tax zone", "in_zone", "text", ex(tp + "tax_zone"))
    s.field("Resident in Bangladesh", "in_resident", "text", "Yes" if ex(tp + "resident", True) else "No", None, YESNO)
    s.field("Gender", "in_gender", "text", {"female": "Female", "third_gender": "Third gender"}.get(ex(tp + "gender"), "Male"),
            "Women get the higher tax-free limit.", ["Male", "Female", "Third gender"])
    s.field("Date of birth", "in_dob", "date", ex(tp + "dob"), "Aged 65 or more on 30 June gets the higher limit.")
    cats = ex(tp + "category", []) or []
    s.field("Person with disability", "in_disabled", "text", "Yes" if "disabled" in cats else "No", None, YESNO)
    s.field("Gazetted war-wounded freedom fighter", "in_ff", "text", "Yes" if "war_wounded_freedom_fighter" in cats else "No", None, YESNO)
    s.field("July Joddha", "in_july", "text", "Yes" if "july_joddha" in cats else "No", None, YESNO)
    s.field("Parent or guardian of a person with disability", "in_parent", "text",
            "Yes" if ex(tp + "parent_of_disabled") else "No", "Adds Tk 50,000 to the tax-free limit.", YESNO)
    s.field("Tick the special-benefit boxes on page 1", "in_boxes", "text",
            "No" if ex("form.special_benefit_boxes", "auto") != "auto" else "Yes",
            "Yes ticks the boxes for the categories above (the sample left them blank).", YESNO)
    s.field("First return", "in_first", "text", "Yes" if ex(tp + "first_return") else "No",
            "Yes: no IT-10B last year; the wealth statement balances on 'Receipt of gift and others'.", YESNO)
    s.field("Father's or husband's name", "in_father", "text", ex(tp + "father_or_husband"))
    s.field("That name is the", "in_relation", "text", "Husband" if ex(tp + "relation") == "husband" else "Father",
            None, ["Father", "Husband"])
    s.field("Spouse's name", "in_spouse", "text", ex(tp + "spouse_name"))
    s.field("Spouse's TIN", "in_spouse_tin", "text", ex(tp + "spouse_tin"))
    s.field("Address", "in_address", "text", ex(tp + "address"))
    s.field("Telephone", "in_tel", "text", ex(tp + "telephone"))
    s.field("Mobile", "in_mobile", "text", ex(tp + "mobile"))
    s.field("Email", "in_email", "text", ex(tp + "email"))
    s.field("Employer", "in_employer", "text", ex(tp + "employer"))
    s.field("Name of organisation (business)", "in_org", "text", ex(tp + "organization"))
    s.field("Business identification number (BIN)", "in_bin", "text", ex(tp + "bin"))
    s.field("Partners and their TINs (firm or AOP)", "in_partners", "text", ex(tp + "partners"))

    # ---------------------------------------------------------------- salary
    s.heading("Salary (Schedule 1)", "Government pay scale: only the items marked taxable on the Rules sheet are taxed.")
    s.sub("Government pay scale")
    labels = ["Basic pay", "Arrear pay", "Special allowance", "House rent allowance", "Medical allowance",
              "Conveyance allowance", "Festival allowance", "Allowance for support staff", "Leave allowance",
              "Honorarium, reward, fee", "Overtime allowance", "Baishakhi allowance", "Interest accrued on provident fund",
              "Lamp grant", "Gratuity", "Others"]
    for k, lab in zip(GOVT_ROWS, labels):
        s.field(lab, f"in_g_{k}", "num", ex(f"salary.govt.{k}"))
    s.sub("Private sector and others")
    labels = ["Basic pay", "Allowances", "Advance salary", "Gratuity, annuity, pension", "Perquisites",
              "In lieu of or in addition to salary", "Employees' share scheme", "Accommodation benefit", "Car benefit",
              "Other benefits from employer", "Employer's contribution to recognised PF", "Others"]
    for k, lab in zip(PRIVATE_ROWS, labels):
        s.field(lab, f"in_p_{k}", "num", ex(f"salary.private.{k}"),
                "Exempt: one third of the total or the cap on the Rules sheet, whichever is lower." if k == "basic_pay" else None)

    # ---------------------------------------------------------------- house property
    s.heading("House property (Schedule 2)", "Up to three properties; they are added together on Schedule 2.")
    for j in range(3):
        c = ws.cell(s.row, 3 + j, f"Property {j + 1}")
        c.font = Font(name="Arial", size=10, bold=True)
    s.row += 1
    hp_rows = [("Location and description", "desc", "text"), ("Use", "use", "text"),
               ("Rent received or annual value, whichever is higher", "av", "num"), ("Advance rent received", "adv", "num"),
               ("Value of other benefits", "ben", "num"), ("Adjusted advance rent", "adj", "num"),
               ("Vacancy allowance", "vac", "num"), ("Municipal or local tax", "mun", "num"), ("Land revenue", "land", "num"),
               ("Interest on loan or mortgage", "loan", "num"), ("Insurance premium", "ins", "num"),
               ("Other deductions", "oth", "num"), ("Your share (%), blank if you own all of it", "share", "num")]
    src = {"desc": "description", "use": "type", "av": "annual_value", "adv": "advance_rent", "ben": "other_benefits",
           "adj": "adjusted_advance_rent", "vac": "vacancy_allowance", "mun": "municipal_tax", "land": "land_revenue",
           "loan": "loan_interest", "ins": "insurance_premium", "oth": "other_deductions", "share": "ownership_share"}
    for lab, key, kind in hp_rows:
        ws.cell(s.row, 2, lab).font = Font(name="Arial", size=10)
        for j in range(3):
            v = ex(f"house_property.{j}.{src[key]}")
            if key == "use":
                v = "Commercial" if v == "commercial" else "Residential"
            if key == "share" and v is not None:
                v = v * 100
            s.cell(3 + j, f"in_hp{j + 1}_{key}", kind, v, ["Residential", "Commercial"] if key == "use" else None)
        if key == "use":
            s.note("Residential: 25% repair allowance; commercial: 30%.")
        s.row += 1

    # ---------------------------------------------------------------- agriculture and business
    s.heading("Agriculture (Schedule 3)")
    s.field("Nature of agriculture", "in_ag_nature", "text", ex("agriculture.nature"))
    s.field("Sale, turnover, receipts", "in_ag_turnover", "num", ex("agriculture.turnover"))
    s.field("Gross profit", "in_ag_gp", "num", ex("agriculture.gross_profit"))
    s.field("Expenses, land revenue, interest, insurance", "in_ag_exp", "num", ex("agriculture.expenses"))
    s.heading("Business or profession (Schedule 4)")
    s.field("Name of business or profession", "in_bz_name", "text", ex("business.name"))
    s.field("Nature of business", "in_bz_nature", "text", ex("business.nature"))
    s.field("Address", "in_bz_address", "text", ex("business.address"))
    s.field("Sale, turnover, receipts", "in_bz_turnover", "num", ex("business.turnover"))
    s.field("Gross profit", "in_bz_gp", "num", ex("business.gross_profit"))
    s.field("General, administrative, selling expenses", "in_bz_exp", "num", ex("business.expenses"))
    s.field("Bad debt", "in_bz_bad", "num", ex("business.bad_debt"))
    s.sub("Balance sheet at 30 June (leave all blank if none)")
    first = s.row
    for lab, key in (("Cash and bank", "cash_bank"), ("Inventory", "inventory"), ("Fixed assets", "fixed_assets"),
                     ("Other assets", "other_assets"), ("Opening capital", "opening_capital"),
                     ("Drawings during the year", "drawings"), ("Liabilities", "liabilities")):
        s.field(lab, f"in_bs_{key}", "num", ex(f"business.balance_sheet.{key}"))
    add_name(wb, "in_bs_all", f"Input!$C${first}:$C${s.row - 1}")

    # ---------------------------------------------------------------- other income
    s.heading("Other income (return lines 5-10)")
    s.field("Capital gains", "in_oi_cg", "num", ex("other_income.capital_gains"))
    s.field("Bank interest, dividend, securities", "in_oi_fin", "num", ex("other_income.financial_assets"))
    s.field("Savings certificate (sanchayapatra) interest", "in_oi_sp", "num", ex("other_income.sanchayapatra_interest"),
            "Taxed finally at source: enter its AIT under tax paid.")
    s.field("Other sources (royalty, honorarium, fees)", "in_oi_other", "num", ex("other_income.other_sources"))
    s.field("Share of income from firm or AOP", "in_oi_firm", "num", ex("other_income.firm_share"))
    s.field("Income of minor or spouse", "in_oi_minor", "num", ex("other_income.minor_spouse"))
    s.field("Foreign income", "in_oi_foreign", "num", ex("other_income.foreign"))

    s.heading("Tax-free income (return line 26)", "The salary exemption is added automatically to row 01.")
    labels = ["Government provident fund (up to Tk 2.5 crore)", "Income already taxed, distributed to beneficiaries",
              "Income of a provident fund (Act 1925)", "Voluntary retirement scheme",
              "Pensioners' savings certificate (up to Tk 5 lakh)", "Share of firm income taxed in the firm",
              "Share from a Hindu undivided family", "Foreign remittance", "Wage earners' development fund",
              "Income from Bandarban, Khagrachari, Rangamati", "Agricultural income (exempt part)",
              "Software and IT business", "Export of handicrafts", "Small and medium industries", "Zero coupon bonds",
              "Government pension", "Welfare allowance, Muktijoddha Kallyan Trust", "Award money from the Government",
              "Management of a briddhashram", "Bangladeshi-flagged ship bringing foreign currency",
              "Other 6th Schedule exemption"]
    s.field("Other salary exemption (row 01)", "in_tf_salary_other", "num", ex("exempt_income.salary_other"))
    for k, lab in zip(TAXFREE_ROWS, labels):
        s.field(lab, f"in_tf_{k}", "num", ex(f"exempt_income.{k}"))

    s.heading("Investments for the rebate (Schedule 5)", "Enter allowable amounts (for example DPS within its limit).")
    labels = ["Life insurance premium", "Deposit pension scheme (DPS)",
              "Government securities, savings certificates, unit certificates", "Listed shares and securities",
              "Provident fund (Act 1925)", "Recognised provident fund (own + employer)", "Approved pension fund",
              "Benevolent fund and group insurance", "Zakat fund", "Others"]
    for k, lab in zip(INVEST_ROWS, labels):
        s.field(lab, f"in_inv_{k}", "num", ex(f"investments.{k}"))

    # ---------------------------------------------------------------- payments
    s.heading("Tax paid (return lines 20-23)")
    s.field("Tax deducted from salary", "in_pay_tds", "num", ex("payments.tds_salary"))
    s.field("AIT on savings certificates", "in_pay_sp", "num", ex("payments.ait_sanchayapatra"))
    s.field("AIT on bank interest", "in_pay_bank", "num", ex("payments.ait_bank_interest"))
    s.field("AIT on WPPF", "in_pay_wppf", "num", ex("payments.ait_wppf"))
    s.field("AIT with car tax token", "in_pay_car", "num", ex("payments.ait_car"))
    s.field("Advance tax by treasury challan", "in_pay_adv", "num", ex("payments.advance_tax"))
    s.field("Refund adjusted from earlier years", "in_pay_refund", "num", ex("payments.refund_adjustment"))
    s.field("Assessment years of that refund", "in_pay_refund_years", "text", ex("payments.refund_assessment_years"))
    s.field("Paid with the return (A-challan)", "in_pay_with_return", "num", ex("payments.paid_with_return"),
            "Leave blank to assume the balance due is paid with the return.")
    s.field("A-challan number", "in_pay_challan", "text", ex("payments.challan_number"))
    s.heading("Other charges and overrides")
    s.field("Owns more than one car, or a house over 8,000 sft in a city corporation", "in_tx_trigger", "text",
            "Yes" if ex("tax.surcharge_trigger") else "No", "Yes charges the 10% surcharge even below Tk 4 crore.", YESNO)
    s.field("Environmental surcharge (cars)", "in_tx_env", "num", ex("tax.environmental_surcharge"))
    s.field("Delay interest or penalty", "in_tx_delay", "num", ex("tax.delay_interest"))
    s.field("Minimum tax u/s 163(2)(kha)", "in_tx_2kha", "num", ex("tax.min_tax_163_2kha"))
    s.field("Minimum tax u/s 163(5)(ka)", "in_tx_5ka", "num", ex("tax.min_tax_163_5ka"))
    s.field("Rebate allowed", "in_tx_rebate", "text", "No" if ex("tax.rebate_allowed") is False else "Yes", None, YESNO)
    s.field("Minimum tax override", "in_tx_min_override", "num", ex("tax.minimum_tax"), "Leave blank to use the Rules sheet.")
    s.field("Net wealth surcharge override", "in_tx_sur_override", "num", ex("tax.net_wealth_surcharge"), "Leave blank to calculate.")
    s.field("Final tax on savings certificates override", "in_tx_final_tax", "num", ex("tax.final_tax"),
            "Leave blank to use the AIT on savings certificates.")

    # ---------------------------------------------------------------- lifestyle
    s.heading("Living expenses (IT-10BB)", "The note column prints in the Comments column of IT-10BB.")
    ws.cell(s.row, 3, "Amount").font = Font(name="Arial", size=10, bold=True)
    ws.cell(s.row, 4, "Note").font = Font(name="Arial", size=10, bold=True)
    s.row += 1
    for lab, key in (("Food, clothing and essentials", "food_clothing"), ("Housing", "housing"),
                     ("Personal transport", "transport"), ("Electricity, gas, water, phone, internet", "utility"),
                     ("Education", "education"), ("Travel and holidays", "travel"),
                     ("Festivals and special occasions", "festival"),
                     ("Tax paid (blank = this return's tax paid)", "tax_paid"), ("Interest on personal loans", "loan_interest")):
        ws.cell(s.row, 2, lab).font = Font(name="Arial", size=10)
        s.cell(3, f"in_lf_{key}", "num", ex(f"lifestyle.{key}"))
        s.cell(4, f"in_lf_{key}_note", "text", ex(f"lifestyle.comments.{key}"))
        s.row += 1

    # ---------------------------------------------------------------- wealth
    s.heading("Assets and liabilities (IT-10B)")
    bal = {"cash": "Cash and bank", "gift": "Gift and other receipts", "none": "None"}.get(ex("wealth.balancing_item"), "Automatic")
    s.field("Net wealth on last year's IT-10B", "in_w_prev", "num", ex("wealth.previous_net_wealth"), "Leave blank on a first return.")
    s.field("Balance the statement with", "in_w_balance", "text", bal,
            "Automatic: gift on a first return, otherwise cash and bank.", ["Automatic", "Cash and bank", "Gift and other receipts", "None"])
    s.field("Receipt of gift and others", "in_w_gift", "num", ex("wealth.gift_and_other_receipts"), "Used unless it is the balancing figure.")
    s.field("Label after 'Receipt of gift and others'", "in_w_gift_label", "text", ex("wealth.gift_label"),
            "Blank prints (First Return) on a first return.")
    s.field("Gifts given, losses, expenses not in IT-10BB", "in_w_other_out", "num", ex("wealth.other_outflows"))
    s.sub("Liabilities")
    s.field("Bank and institutional loans", "in_w_liab_inst", "num", ex("wealth.liabilities.institutional"))
    s.field("Non-institutional loans", "in_w_liab_non", "num", ex("wealth.liabilities.non_institutional"))
    s.field("Other liabilities", "in_w_liab_other", "num", ex("wealth.liabilities.other"))
    s.sub("Non-agricultural property (up to 4)")
    ws.cell(s.row, 3, "Description, location and size").font = Font(name="Arial", size=10, bold=True)
    ws.cell(s.row, 5, "Cost value").font = Font(name="Arial", size=10, bold=True)
    s.row += 1
    for j in range(4):
        ws.cell(s.row, 2, f"Property {j + 1}").font = Font(name="Arial", size=10)
        s.cell(3, f"in_na{j + 1}_desc", "text", ex(f"wealth.assets.non_agricultural_property.{j}.description"))
        ws.merge_cells(start_row=s.row, start_column=3, end_row=s.row, end_column=4)
        s.cell(5, f"in_na{j + 1}_val", "num", ex(f"wealth.assets.non_agricultural_property.{j}.value"))
        s.row += 1
    s.sub("Agricultural land (up to 2)")
    for j in range(2):
        ws.cell(s.row, 2, f"Land {j + 1}").font = Font(name="Arial", size=10)
        s.cell(3, f"in_agp{j + 1}_desc", "text", ex(f"wealth.assets.agricultural_property.{j}.description"))
        ws.merge_cells(start_row=s.row, start_column=3, end_row=s.row, end_column=4)
        s.cell(5, f"in_agp{j + 1}_val", "num", ex(f"wealth.assets.agricultural_property.{j}.value"))
        s.row += 1
    s.sub("Financial assets")
    for lab, nm, key in (("Debentures, bonds, shares, unit certificates", "deb", "debenture_bond"),
                         ("Savings certificates and DPS", "sp", "sanchayapatra_dps"),
                         ("Life insurance (premiums paid)", "li", "life_insurance"), ("Loans given", "loan", "loan_given"),
                         ("Fixed deposits", "fdr", "fdr"), ("Provident fund and other funds", "pf", "provident_fund"),
                         ("Other investments", "oinv", "other_investment")):
        s.field(lab, f"in_w_{nm}", "num", ex(f"wealth.assets.{key}"))
    s.sub("Other assets")
    s.field("Business capital (only if Schedule 4 has no balance sheet)", "in_w_bcap_agg", "num", ex("wealth.assets.business_capital_aggregate"))
    s.field("Business liabilities (same)", "in_w_bcap_liab", "num", ex("wealth.assets.business_liabilities"))
    s.field("Director's shareholdings in limited companies", "in_w_dir", "num", ex("wealth.assets.director_shareholdings"))
    s.field("Motor car (cost)", "in_w_car", "num", ex("wealth.assets.motor_car"))
    s.field("Car type and registration number", "in_w_car_det", "text", ex("wealth.assets.motor_car_details"))
    s.field("Ornaments (value)", "in_w_orn", "num", ex("wealth.assets.ornaments"))
    s.field("Ornaments quantity", "in_w_orn_qty", "text", ex("wealth.assets.ornaments_quantity"), "For example: 25 Bhori")
    s.field("Ornaments type", "in_w_orn_type", "text", ex("wealth.assets.ornaments_type", "(Gold)"))
    s.field("Ornaments note", "in_w_orn_note", "text", ex("wealth.assets.ornaments_note"), "For example: (Wedding Gifts)")
    s.field("Furniture and electronics", "in_w_furn", "num", ex("wealth.assets.furniture_electronics"))
    s.field("Furniture note", "in_w_furn_note", "text", ex("wealth.assets.furniture_note"))
    s.field("Other assets", "in_w_oth", "num", ex("wealth.assets.other_assets"))
    s.field("Cash, bank and funds outside business", "in_w_cash", "num", ex("wealth.assets.cash_bank"),
            "Calculated instead when cash and bank is the balancing figure.")

    # ---------------------------------------------------------------- signing
    s.heading("Signing and documents")
    s.field("Place", "in_place", "text", ex("form.place"))
    s.field("Date", "in_date", "date", ex("form.date"))
    docs = ex("form.documents") or ["Tax Calculation", "Tax Free Income/ Exempted", "Related Supporting Documents"]
    for j in range(7):
        s.field(f"Document {j + 1:02d}", f"in_doc{j + 1}", "text", docs[j] if j < len(docs) else None)
    ws.print_area = f"A1:F{s.row}"
    return ws


# =================================================================== Rules sheet
def build_rules(wb):
    ws = wb.create_sheet("Rules")
    print_wide(ws, landscape=True)
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 2), ("B", 30), ("C", 54), ("D", 17), ("E", 17), ("F", 17), ("G", 15)):
        ws.column_dimensions[col].width = w * CHARS
    ws["B1"] = "Tax rules"
    ws["B1"].font = Font(name="Times New Roman", size=18, bold=True, color=GREEN)
    ws["B2"] = ("Edit the blue figures if the official Finance Act differs. Column G follows the 'Tax rules' choice on the "
                "Input sheet. AY 2026-27 figures come from news reports of the Finance Act 2026, not the gazette.")
    ws["B2"].font = Font(name="Arial", size=10)
    ws["B4"], ws["C4"] = "Parameter", "Meaning"
    for j, p in enumerate(PROFILES):
        ws.cell(4, 4 + j, p)
    ws["G4"] = "In use"
    for c in ws[4]:
        c.font = Font(name="Arial", size=10, bold=True)
    add_name(wb, "rules_ids", "Rules!$D$4:$F$4")

    rows = []
    T = lambda p: RULES[p]["thresholds"]  # noqa: E731
    for key, lab in (("general", "Tax-free limit: general"), ("female", "Tax-free limit: women"),
                     ("senior", "Tax-free limit: aged 65 or more"), ("disabled", "Tax-free limit: persons with disability"),
                     ("third_gender", "Tax-free limit: third gender"),
                     ("war_wounded_freedom_fighter", "Tax-free limit: gazetted war-wounded freedom fighter"),
                     ("july_joddha", "Tax-free limit: July Joddha"),
                     ("parent_of_disabled_addition", "Added for a parent or guardian of a person with disability")):
        rows.append((f"thr_{key}", lab, [T(p)[key] for p in PROFILES], "#,##0"))
    for i in range(5):
        if i < 4:
            rows.append((f"w{i + 1}", f"Slab {i + 2} width (after the tax-free limit)",
                         [RULES[p]["slabs"][i]["width"] for p in PROFILES], "#,##0"))
        rows.append((f"r{i + 1}", f"Slab {i + 2} rate" + (" (rest of the income)" if i == 4 else ""),
                     [RULES[p]["slabs"][i]["rate"] for p in PROFILES], "0%"))
    rows += [
        ("min_tax", "Minimum tax", [RULES[p]["minimum_tax"]["amount"] for p in PROFILES], "#,##0"),
        ("min_tax_first", "Minimum tax on a first return", [RULES[p]["minimum_tax"]["first_return_amount"] for p in PROFILES], "#,##0"),
        ("rebate_income_rate", "Rebate limit: share of taxable income", [RULES[p]["rebate"]["income_rate"] for p in PROFILES], "0%"),
        ("rebate_invest_rate", "Rebate limit: share of eligible investment", [RULES[p]["rebate"]["investment_rate"] for p in PROFILES], "0%"),
        ("rebate_cap", "Rebate limit: maximum", [RULES[p]["rebate"]["cap"] for p in PROFILES], "#,##0"),
        ("priv_fraction", "Private salary exemption: share of salary", [RULES[p]["private_salary_exemption"]["fraction"] for p in PROFILES], "0.0000"),
        ("priv_cap", "Private salary exemption: maximum", [RULES[p]["private_salary_exemption"]["cap"] for p in PROFILES], "#,##0"),
        ("repair_res", "Repair allowance: residential property", [RULES[p]["house_repair_rate"]["residential"] for p in PROFILES], "0%"),
        ("repair_com", "Repair allowance: commercial property", [RULES[p]["house_repair_rate"]["commercial"] for p in PROFILES], "0%"),
    ]
    for i in range(4):
        rows.append((f"sur_up{i + 1}", f"Surcharge band {i + 1}: net wealth up to",
                     [RULES[p]["surcharge"]["bands"][i]["upto"] for p in PROFILES], "#,##0"))
    for i in range(5):
        rows.append((f"sur_rate{i + 1}", f"Surcharge band {i + 1}: rate" + (" (above the last limit)" if i == 4 else ""),
                     [RULES[p]["surcharge"]["bands"][i]["rate"] for p in PROFILES], "0%"))
    rows += [
        ("sur_base", "Surcharge on: 1 = tax after rebate, 0 = tax payable incl. minimum tax",
         [1 if RULES[p]["surcharge"]["base"] == "regular_tax" else 0 for p in PROFILES], "0"),
        ("sur_trigger", "Surcharge when more than one car or a large city house", [RULES[p]["surcharge"]["trigger_rate"] for p in PROFILES], "0%"),
    ]
    labels = ["Basic pay", "Arrear pay", "Special allowance", "House rent allowance", "Medical allowance",
              "Conveyance allowance", "Festival allowance", "Allowance for support staff", "Leave allowance",
              "Honorarium, reward, fee", "Overtime allowance", "Baishakhi allowance", "Interest accrued on provident fund",
              "Lamp grant", "Gratuity", "Others"]
    for k, lab in zip(GOVT_ROWS, labels):
        rows.append((f"gt_{k}", f"Government pay scale: {lab} taxable (1 = yes)",
                     [1 if k in RULES[p]["govt_salary_taxable_items"] else 0 for p in PROFILES], "0"))

    r = 5
    for key, lab, vals, fmt in rows:
        ws.cell(r, 2, key).font = Font(name="Arial", size=9, color="5D6B64")
        ws.cell(r, 3, lab).font = Font(name="Arial", size=10)
        for j, v in enumerate(vals):
            c = ws.cell(r, 4 + j, v)
            c.font = Font(name="Arial", size=10, color=BLUE)
            c.number_format = fmt
        g = ws.cell(r, 7, f"=INDEX(D{r}:F{r},MATCH(in_profile,rules_ids,0))")
        g.number_format = fmt
        g.font = Font(name="Arial", size=10, bold=True)
        add_name(wb, f"p_{key}", f"Rules!$G${r}")
        r += 1
    r += 1
    ws.cell(r, 2, "notes").font = Font(name="Arial", size=9, color="5D6B64")
    ws.cell(r, 3, "Caution shown on the Input sheet").font = Font(name="Arial", size=10)
    for j, p in enumerate(PROFILES):
        c = ws.cell(r, 4 + j, RULES[p].get("warning", ""))
        c.font = Font(name="Arial", size=9, color=BLUE)
    ws.cell(r, 7, f"=INDEX(D{r}:F{r},MATCH(in_profile,rules_ids,0))")
    add_name(wb, "p_warning", f"Rules!$G${r}")
    r += 2
    ws.cell(r, 2, "Sources").font = Font(name="Arial", size=10, bold=True)
    for line in ("The Daily Star, 29 June 2026: Finance bill passed, tax-free income threshold set at Tk 4 lakh for FY2026-27.",
                 "bdnews24: New income tax slabs in effect from July, says NBR chief.",
                 "The Financial Express: Minimum income tax Tk 5,000, Tk 1,000 for first-time taxpayers.",
                 "The Daily Star: How Bangladesh's surcharge and wealth tax system works.",
                 "Sample template: rules used by the sample return (older 5% slab, women Tk 4,00,000)."):
        r += 1
        ws.cell(r, 3, line).font = Font(name="Arial", size=9)
    return ws


# =================================================================== Calc sheet
def calc_rows():
    """(key, description, formula) in order. Keys match the renderer's field keys."""
    R = []

    def add(key, desc, f):
        R.append((key, desc, f))

    add("ay", "Assessment year", '=in_ay&""')
    add("ay_start", "First year of the assessment year", "=VALUE(LEFT(in_ay,4))")
    add("iy_end_mdy", "Income year end (m/d/yyyy)", '="6/30/"&v_ay_start')
    add("iy_end_dmy", "Income year end (d-Mon-yy)", '="30-Jun-"&RIGHT(v_ay_start,2)')
    add("name", "Name", '=TRIM(in_name&"")')
    add("name_line", "Name line", '="Name: "&v_name')
    add("tin", "TIN digits", '=SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM(in_tin&"")," ",""),"-",""),".","")')
    for k, src in (("nid", "in_nid"), ("circle", "in_circle"), ("zone", "in_zone"), ("spouse_name", "in_spouse"),
                   ("spouse_tin", "in_spouse_tin"), ("address", "in_address"), ("telephone", "in_tel"),
                   ("mobile", "in_mobile"), ("email", "in_email"), ("employer", "in_employer"),
                   ("organization", "in_org"), ("bin", "in_bin"), ("partners", "in_partners"),
                   ("father_or_husband", "in_father"), ("place", "in_place")):
        add(k, k.replace("_", " ").capitalize(), f'={src}&""')
    add("resident", "Resident", '=in_resident<>"No"')
    add("nonresident", "Non-resident", '=in_resident="No"')
    add("dob_digits", "Date of birth as DDMMYYYY", '=IF(N(in_dob)=0,"",TEXT(DAY(in_dob),"00")&TEXT(MONTH(in_dob),"00")&YEAR(in_dob))')
    add("age", "Age on 30 June", '=IF(N(in_dob)=0,"",DATEDIF(in_dob,DATE(v_ay_start,6,30),"y"))')
    add("is_female", "Woman", '=in_gender="Female"')
    add("is_third", "Third gender", '=in_gender="Third gender"')
    add("is_senior", "Aged 65 or more", "=N(v_age)>=65")
    add("is_disabled", "Person with disability", '=in_disabled="Yes"')
    add("is_ff", "Gazetted war-wounded freedom fighter", '=in_ff="Yes"')
    add("is_july", "July Joddha", '=in_july="Yes"')
    add("is_parent", "Parent of a person with disability", '=in_parent="Yes"')
    add("first_return", "First return", '=in_first="Yes"')
    add("threshold", "Tax-free limit (highest that applies)",
        "=MAX(p_thr_general,IF(v_is_female,p_thr_female,0),IF(v_is_senior,p_thr_senior,0),IF(v_is_disabled,p_thr_disabled,0),"
        "IF(v_is_third,p_thr_third_gender,0),IF(v_is_ff,p_thr_war_wounded_freedom_fighter,0),IF(v_is_july,p_thr_july_joddha,0))"
        "+IF(v_is_parent,p_thr_parent_of_disabled_addition,0)")
    add("boxes_auto", "Tick special-benefit boxes", '=in_boxes<>"No"')
    add("box.ff", "Box: freedom fighter / July Joddha", "=AND(v_boxes_auto,OR(v_is_ff,v_is_july))")
    add("box.female", "Box: female", "=AND(v_boxes_auto,v_is_female)")
    add("box.third", "Box: third gender", "=AND(v_boxes_auto,v_is_third)")
    add("box.disabled", "Box: disabled", "=AND(v_boxes_auto,v_is_disabled)")
    add("box.65", "Box: aged 65 or more", "=AND(v_boxes_auto,v_is_senior)")
    add("box.parent", "Box: parent of a person with disability", "=AND(v_boxes_auto,v_is_parent)")
    add("relation", "Father or husband", '=IF(in_relation="Husband","husband","father")')
    add("date", "Date of signing", '=IF(N(in_date)=0,"",TEXT(DAY(in_date),"00")&"/"&TEXT(MONTH(in_date),"00")&"/"&YEAR(in_date))')
    for j in range(1, 8):
        add(f"doc{j}", f"Document {j}", f'=in_doc{j}&""')

    # Schedule 1 (a): incomes, then exempt column, then taxable column (contiguous for SUM).
    for k in GOVT_ROWS:
        add(f"s1a.{k}.inc", f"Govt pay scale: {k} income", f"=N(in_g_{k})")
    for k in GOVT_ROWS:
        add(f"s1a.{k}.exm", f"Govt pay scale: {k} exempt", f'=IF(p_gt_{k}=1,"",v_s1a_{k}_inc)')
    for k in GOVT_ROWS:
        add(f"s1a.{k}.tax", f"Govt pay scale: {k} taxable", f'=IF(p_gt_{k}=1,v_s1a_{k}_inc,"")')
    add("s1a.inc", "Govt pay scale: total income", "=SUM({s1a.%s.inc}:{s1a.%s.inc})" % (GOVT_ROWS[0], GOVT_ROWS[-1]))
    add("s1a.exm", "Govt pay scale: total exempt", "=SUM({s1a.%s.exm}:{s1a.%s.exm})" % (GOVT_ROWS[0], GOVT_ROWS[-1]))
    add("s1a.tax", "Govt pay scale: total taxable", "=SUM({s1a.%s.tax}:{s1a.%s.tax})" % (GOVT_ROWS[0], GOVT_ROWS[-1]))
    for k in PRIVATE_ROWS:
        add(f"s1b.{k}", f"Private salary: {k}", f"=N(in_p_{k})")
    add("s1b.sub", "Private salary: total", "=SUM({s1b.%s}:{s1b.%s})" % (PRIVATE_ROWS[0], PRIVATE_ROWS[-1]))
    add("s1b.total", "Private salary: total (line 13)", "=v_s1b_sub")
    add("s1b.exempt", "Private salary: exempt (1/3 or cap)", "=IF(v_s1b_total>0,ROUND(MIN(v_s1b_total*p_priv_fraction,p_priv_cap),0),0)")
    add("s1b.net", "Private salary: taxable", "=v_s1b_total-v_s1b_exempt")

    # Schedule 2, per property then added up.
    for j in (1, 2, 3):
        p = f"in_hp{j}_"
        add(f"hp{j}.r1", f"Property {j}: annual value", f"=N({p}av)")
        add(f"hp{j}.r2", f"Property {j}: advance rent", f"=N({p}adv)")
        add(f"hp{j}.r3", f"Property {j}: other benefits", f"=N({p}ben)")
        add(f"hp{j}.r4", f"Property {j}: adjusted advance rent", f"=N({p}adj)")
        add(f"hp{j}.r5", f"Property {j}: vacancy allowance", f"=N({p}vac)")
        add(f"hp{j}.r6", f"Property {j}: total rent income", f"=v_hp{j}_r1+v_hp{j}_r2+v_hp{j}_r3-v_hp{j}_r4-v_hp{j}_r5")
        add(f"hp{j}.com", f"Property {j}: commercial", f'={p}use="Commercial"')
        add(f"hp{j}.da", f"Property {j}: repair allowance", f"=ROUND(v_hp{j}_r6*IF(v_hp{j}_com,p_repair_com,p_repair_res),0)")
        add(f"hp{j}.db", f"Property {j}: municipal tax", f"=N({p}mun)")
        add(f"hp{j}.dc", f"Property {j}: land revenue", f"=N({p}land)")
        add(f"hp{j}.dd", f"Property {j}: loan interest", f"=N({p}loan)")
        add(f"hp{j}.de", f"Property {j}: insurance premium", f"=N({p}ins)")
        add(f"hp{j}.df", f"Property {j}: other deductions", f"=N({p}oth)")
        add(f"hp{j}.r8", f"Property {j}: total deductions",
            f"=v_hp{j}_da+v_hp{j}_db+v_hp{j}_dc+v_hp{j}_dd+v_hp{j}_de+v_hp{j}_df")
        add(f"hp{j}.r9", f"Property {j}: income", f"=v_hp{j}_r6-v_hp{j}_r8")
        add(f"hp{j}.share", f"Property {j}: ownership share", f"=IF(OR(N({p}share)<=0,N({p}share)>=100),1,N({p}share)/100)")
        add(f"hp{j}.income", f"Property {j}: your income", f"=IF(v_hp{j}_share<1,ROUND(v_hp{j}_r9*v_hp{j}_share,0),v_hp{j}_r9)")
        add(f"hp{j}.used", f"Property {j}: filled in",
            f'=OR({p}desc&""<>"",v_hp{j}_r1<>0,v_hp{j}_r2<>0,v_hp{j}_r3<>0,v_hp{j}_r4<>0,v_hp{j}_r5<>0,v_hp{j}_r8<>0)')
    for k in ("r1", "r2", "r3", "r4", "r5", "r6", "da", "db", "dc", "dd", "de", "df", "r8", "r9"):
        add(f"s2.{k}", f"Schedule 2 {k}", f"=v_hp1_{k}+v_hp2_{k}+v_hp3_{k}")
    add("s2.r7", "Schedule 2 deductions (line 7 total)", "=v_s2_r8")
    add("s2.r10", "Schedule 2 share of income (part ownership)",
        '=IF(OR(v_hp1_share<1,v_hp2_share<1,v_hp3_share<1),v_hp1_income+v_hp2_income+v_hp3_income,"")')
    add("s2.nused", "Properties filled in", "=v_hp1_used*1+v_hp2_used*1+v_hp3_used*1")
    add("s2.ncom", "Commercial properties", "=AND(v_hp1_used,v_hp1_com)*1+AND(v_hp2_used,v_hp2_com)*1+AND(v_hp3_used,v_hp3_com)*1")
    add("s2.rate", "Repair allowance shown",
        '=IF(v_s2_ncom=0,TEXT(p_repair_res*100,"0")&"%",IF(v_s2_ncom=v_s2_nused,TEXT(p_repair_com*100,"0")&"%","Mixed"))')
    add("s2.desc", "Property descriptions",
        '=in_hp1_desc&IF(AND(in_hp1_desc&""<>"",in_hp2_desc&""<>""),"; ","")&in_hp2_desc'
        '&IF(AND(in_hp1_desc&in_hp2_desc<>"",in_hp3_desc&""<>""),"; ","")&in_hp3_desc')
    add("house_income", "Income from house property", "=v_hp1_income*v_hp1_used+v_hp2_income*v_hp2_used+v_hp3_income*v_hp3_used")

    add("s3.nature", "Nature of agriculture", '=in_ag_nature&""')
    add("s3.r1", "Agriculture turnover", "=N(in_ag_turnover)")
    add("s3.r2", "Agriculture gross profit", "=N(in_ag_gp)")
    add("s3.r3", "Agriculture expenses", "=N(in_ag_exp)")
    add("s3.r4", "Agriculture net profit", "=v_s3_r2-v_s3_r3")
    add("s4.name", "Business name", '=in_bz_name&""')
    add("s4.nature", "Business nature", '=in_bz_nature&""')
    add("s4.address", "Business address", '=in_bz_address&""')
    add("s4.r1", "Turnover", "=N(in_bz_turnover)")
    add("s4.r2", "Gross profit", "=N(in_bz_gp)")
    add("s4.r3", "Expenses", "=N(in_bz_exp)")
    add("s4.r4", "Bad debt", "=N(in_bz_bad)")
    add("s4.r5", "Net profit", "=v_s4_r2-v_s4_r3-v_s4_r4")
    add("s4.r6", "Cash and bank", "=N(in_bs_cash_bank)")
    add("s4.r7", "Inventory", "=N(in_bs_inventory)")
    add("s4.r8", "Fixed assets", "=N(in_bs_fixed_assets)")
    add("s4.r9", "Other assets", "=N(in_bs_other_assets)")
    add("s4.r10", "Total assets", "=v_s4_r6+v_s4_r7+v_s4_r8+v_s4_r9")
    add("s4.r11", "Opening capital", "=N(in_bs_opening_capital)")
    add("s4.r12", "Net profit", "=v_s4_r5")
    add("s4.r13", "Drawings", "=N(in_bs_drawings)")
    add("s4.r14", "Closing capital", "=v_s4_r11+v_s4_r12-v_s4_r13")
    add("s4.r15", "Liabilities", "=N(in_bs_liabilities)")
    add("s4.r16", "Total capital and liabilities", "=v_s4_r14+v_s4_r15")
    add("biz_present", "Business balance sheet entered", "=COUNTA(in_bs_all)>0")

    add("r1", "Line 1 salaries", "=v_s1a_tax+v_s1b_net")
    add("r2", "Line 2 house property", "=v_house_income")
    add("r3", "Line 3 agriculture", "=v_s3_r4")
    add("r4", "Line 4 business", "=v_s4_r5")
    add("r5", "Line 5 capital gains", "=N(in_oi_cg)")
    add("r6", "Line 6 financial assets (incl. savings certificates)", "=N(in_oi_fin)+N(in_oi_sp)")
    add("r7", "Line 7 other sources", "=N(in_oi_other)")
    add("r8", "Line 8 share of firm or AOP", "=N(in_oi_firm)")
    add("r9", "Line 9 minor or spouse", "=N(in_oi_minor)")
    add("r10", "Line 10 foreign income", "=N(in_oi_foreign)")
    add("r11", "Line 11 total income", "=SUM({r1}:{r10})")

    add("tf.r01", "Tax-free row 01: salary exemption", "=v_s1a_exm+v_s1b_exempt+N(in_tf_salary_other)")
    for i, k in enumerate(TAXFREE_ROWS, start=2):
        add(f"tf.r{i:02d}", f"Tax-free row {i:02d}: {k}", f"=N(in_tf_{k})")
    add("tf.total", "Total tax-free income", "=SUM({tf.r01}:{tf.r22})")
    add("r26", "Line 26 tax-exempt income", "=v_tf_total")

    add("tc.A", "A. Total income", "=v_r11")
    add("tc.sp", "Less savings-certificate interest (final tax)", "=N(in_oi_sp)")
    add("tc.An", "Income for the slabs", "=v_tc_A-v_tc_sp")
    add("tc.s0.inc", "First slab income", "=MIN(MAX(v_tc_An,0),v_threshold)")
    add("tc.s1.inc", "Slab 2 income", "=MIN(MAX(v_tc_An-v_threshold,0),p_w1)")
    add("tc.s2.inc", "Slab 3 income", "=MIN(MAX(v_tc_An-v_threshold-p_w1,0),p_w2)")
    add("tc.s3.inc", "Slab 4 income", "=MIN(MAX(v_tc_An-v_threshold-p_w1-p_w2,0),p_w3)")
    add("tc.s4.inc", "Slab 5 income", "=MIN(MAX(v_tc_An-v_threshold-p_w1-p_w2-p_w3,0),p_w4)")
    add("tc.s5.inc", "Rest of the income", "=MAX(v_tc_An-v_threshold-p_w1-p_w2-p_w3-p_w4,0)")
    add("tc.s0.rate", "First slab rate", '="0%"')
    for i in range(1, 6):
        add(f"tc.s{i}.rate", f"Slab {i + 1} rate", f'=TEXT(p_r{i}*100,"0")&"%"')
    add("tc.s0.tax", "First slab tax", "=0")
    for i in range(1, 6):
        add(f"tc.s{i}.tax", f"Slab {i + 1} tax", f"=v_tc_s{i}_inc*p_r{i}")
    add("tc.sinc", "Total slab income", "=MAX(v_tc_An,0)")
    add("tc.B", "B. Tax on slabs (line 12)", "=ROUND(SUM({tc.s0.tax}:{tc.s5.tax}),0)")
    for i, k in enumerate(INVEST_ROWS, start=1):
        add(f"s5.r{i}", f"Investment: {k}", f"=N(in_inv_{k})")
    add("s5.r11", "Total investment", "=SUM({s5.r1}:{s5.r10})")
    add("tc.ci", "Rebate limit i: share of income", "=MAX(v_tc_An,0)*p_rebate_income_rate")
    add("tc.cii", "Rebate limit ii: share of investment", "=v_s5_r11*p_rebate_invest_rate")
    add("tc.ciii", "Rebate limit iii: maximum", "=p_rebate_cap")
    add("tc.C", "C. Rebate (lowest of i, ii, iii)", '=IF(in_tx_rebate="No",0,ROUND(MIN(v_tc_ci,v_tc_cii,v_tc_ciii),0))')
    add("s5.r12", "Amount of tax rebate", "=v_tc_C")
    add("c_used", "Rebate used (not more than the slab tax)", "=MIN(v_tc_C,v_tc_B)")
    add("tc.D", "D. Tax after rebate", "=v_tc_B-v_c_used")
    add("tc.ei", "E i. Minimum tax",
        '=IF(in_tx_min_override&""<>"",N(in_tx_min_override),IF(v_tc_An<=v_threshold,0,IF(v_first_return,p_min_tax_first,p_min_tax)))')
    add("tc.eii", "E ii. Minimum tax u/s 163(2)(kha)", "=N(in_tx_2kha)")
    add("tc.eiii", "E iii. Minimum tax u/s 163(5)(ka)", "=N(in_tx_5ka)")
    add("tc.E", "E. Minimum tax (highest)", "=MAX(v_tc_ei,v_tc_eii,v_tc_eiii)")
    add("tc.F", "F. Tax payable (higher of D and E)", "=MAX(v_tc_D,v_tc_E)")
    for i, (src, lab) in enumerate((("in_pay_tds", "Remuneration (TDS)"), ("in_pay_sp", "AIT on savings certificates"),
                                    ("in_pay_wppf", "AIT on WPPF"), ("in_pay_car", "AIT with car"),
                                    ("in_pay_bank", "AIT on bank interest"), ("in_pay_refund", "Refund adjusted"),
                                    ("in_pay_adv", "Advance tax")), start=1):
        add(f"tc.j{i}", f"J. {lab}", f"=N({src})")
    add("tc.jsub1", "Tax deducted or collected at source", "=v_tc_j1+v_tc_j2+v_tc_j3+v_tc_j4+v_tc_j5")
    add("tc.jsub2", "Refund adjusted and advance tax", "=v_tc_j6+v_tc_j7")
    add("tc.H", "H. Final tax on savings certificates",
        '=IF(in_tx_final_tax&""<>"",N(in_tx_final_tax),IF(N(in_oi_sp)>0,v_tc_j2,0))')
    add("env", "Environmental surcharge", "=N(in_tx_env)")
    add("delay", "Delay interest and penalty", "=N(in_tx_delay)")
    add("paid_given", "Amount paid with the return entered", '=in_pay_with_return&""<>""')

    # Wealth inputs.
    add("a.agg", "Aggregate business capital and liabilities", "=IF(v_biz_present,v_s4_r10,N(in_w_bcap_agg))")
    add("a.bliab", "Business liabilities", "=IF(v_biz_present,v_s4_r15,N(in_w_bcap_liab))")
    add("a.bcap", "Business capital", "=v_a_agg-v_a_bliab")
    add("a.dir", "Director's shareholdings", "=N(in_w_dir)")
    add("a.bcap_other", "Business capital other than director's shareholdings", "=v_a_bcap")
    for j in range(1, 5):
        add(f"a.d{j}_desc", f"Non-agricultural property {j}", f'=in_na{j}_desc&""')
        add(f"a.d{j}_val", f"Non-agricultural property {j} value", f'=IF(AND(in_na{j}_desc&""="",in_na{j}_val&""=""),"",N(in_na{j}_val))')
    add("a.d_total", "Non-agricultural property total", "=N(in_na1_val)+N(in_na2_val)+N(in_na3_val)+N(in_na4_val)")
    add("a.e_total", "Agricultural property total", "=N(in_agp1_val)+N(in_agp2_val)")
    add("a.e_desc", "Agricultural property description",
        '=in_agp1_desc&IF(AND(in_agp1_desc&""<>"",in_agp2_desc&""<>""),"; ","")&in_agp2_desc')
    for i, src in enumerate(("in_w_deb", "in_w_sp", "in_w_li", "in_w_loan", "in_w_fdr", "in_w_pf", "in_w_oinv"), start=1):
        add(f"a.f{i}", f"Financial asset {i}", f"=N({src})")
    add("a.f_total", "Total financial assets", "=SUM({a.f1}:{a.f7})")
    add("a.g", "Motor car", "=N(in_w_car)")
    add("a.g_desc", "Motor car details", '=in_w_car_det&""')
    add("a.h", "Ornaments", "=N(in_w_orn)")
    add("a.h_qty", "Ornaments quantity", '=in_w_orn_qty&""')
    add("a.h_type", "Ornaments type", '=in_w_orn_type&""')
    add("a.h_type_shown", "Ornaments type as printed", '=IF(OR(v_a_h_qty<>"",v_a_h<>0),v_a_h_type,"")')
    add("a.h_note", "Ornaments note", '=in_w_orn_note&""')
    add("a.i", "Furniture and electronics", "=N(in_w_furn)")
    add("a.i_note", "Furniture note", '=in_w_furn_note&""')
    add("a.j", "Other assets", "=N(in_w_oth)")
    add("non_cash", "Assets other than cash",
        "=v_a_bcap+v_a_dir+v_a_d_total+v_a_e_total+v_a_f_total+v_a_g+v_a_h+v_a_i+v_a_j")
    add("w.6a", "Institutional liabilities", "=N(in_w_liab_inst)")
    add("w.6b", "Non-institutional liabilities", "=N(in_w_liab_non)")
    add("w.6c", "Other liabilities", "=N(in_w_liab_other)")
    add("w.liab", "Total liabilities", "=v_w_6a+v_w_6b+v_w_6c")
    add("other_out", "Gifts given, losses, other outflows", "=N(in_w_other_out)")
    add("gift_in", "Receipt of gift and others (entered)", "=N(in_w_gift)")
    add("cash_in", "Cash and bank (entered)", "=N(in_w_cash)")
    add("prev_given", "Last year's net wealth entered", '=in_w_prev&""<>""')
    add("prev", "Last year's net wealth", "=N(in_w_prev)")
    add("mode", "Balancing figure",
        '=IF(in_w_balance="Cash and bank","cash",IF(in_w_balance="Gift and other receipts","gift",IF(in_w_balance="None","none",IF(v_prev_given,"cash","gift"))))')
    add("life_base", "Living expenses other than tax",
        "=N(in_lf_food_clothing)+N(in_lf_housing)+N(in_lf_transport)+N(in_lf_utility)+N(in_lf_education)+N(in_lf_travel)+N(in_lf_festival)+N(in_lf_loan_interest)")

    # Surcharge: net wealth before the surcharge decides the band (a surcharge-free fixed point).
    add("k0", "Balance before surcharge", "=v_tc_F+v_env+v_delay+v_tc_H-v_tc_jsub1-v_tc_jsub2")
    add("p24_0", "Tax paid before surcharge", "=v_tc_jsub1+v_tc_j7+v_tc_j6+IF(v_paid_given,N(in_pay_with_return),MAX(v_k0,0))")
    add("nw0", "Net wealth before surcharge",
        '=IF(v_mode="gift",v_non_cash+v_cash_in-v_w_liab,v_r11+v_r26+v_gift_in+v_prev'
        '-(v_life_base+IF(in_lf_tax_paid&""<>"",N(in_lf_tax_paid),v_p24_0)+v_other_out))')
    add("sur_band_rate", "Surcharge rate from net wealth",
        "=IF(v_nw0<=p_sur_up1,p_sur_rate1,IF(v_nw0<=p_sur_up2,p_sur_rate2,IF(v_nw0<=p_sur_up3,p_sur_rate3,IF(v_nw0<=p_sur_up4,p_sur_rate4,p_sur_rate5))))")
    add("sur_rate", "Surcharge rate", '=IF(AND(in_tx_trigger="Yes",v_sur_band_rate<p_sur_trigger),p_sur_trigger,v_sur_band_rate)')
    add("surcharge", "Net wealth surcharge",
        '=IF(in_tx_sur_override&""<>"",N(in_tx_sur_override),ROUND(IF(p_sur_base=1,v_tc_D,v_tc_F)*v_sur_rate,0))')
    add("tc.gs", "G. Surcharge", "=v_surcharge+v_env")
    add("tc.gd", "G. Delay interest, penalty", "=v_delay")
    add("tc.G", "G. Total", "=v_surcharge+v_env+v_delay")
    add("tc.I", "I. Total payable", "=v_tc_F+v_tc_G+v_tc_H")
    add("tc.K", "K. Balance payable/(refundable)", "=v_tc_I-v_tc_jsub1-v_tc_jsub2")
    add("tc.paid", "Paid with the return by A-challan", "=IF(v_paid_given,N(in_pay_with_return),0)")
    add("tc.challan", "A-challan number", '=in_pay_challan&""')
    add("tc.net", "Net tax payable/(refundable)", "=v_tc_K-v_tc_paid")

    add("r12", "Line 12 gross tax", "=v_tc_B")
    add("r13", "Line 13 rebate", "=v_c_used")
    add("r14", "Line 14 net tax after rebate", "=v_tc_D")
    add("r15", "Line 15 minimum tax", "=v_tc_E")
    add("r16", "Line 16 tax payable", "=v_tc_F")
    add("r17a", "Line 17(a) net wealth surcharge", "=v_surcharge")
    add("r17b", "Line 17(b) environmental surcharge", "=v_env")
    add("r17", "Line 17 surcharges", "=v_r17a+v_r17b")
    add("r18", "Line 18 delay interest, penalty, final tax", "=v_delay+v_tc_H")
    add("r19", "Line 19 total amount payable", "=v_r16+v_r17+v_r18")
    add("r20", "Line 20 tax deducted at source", "=v_tc_jsub1")
    add("r21", "Line 21 advance tax", "=v_tc_j7")
    add("r22", "Line 22 refund adjusted", "=v_tc_j6")
    add("r22_years", "Line 22 assessment years of the refund", '=in_pay_refund_years&""')
    add("r23", "Line 23 paid with the return", "=IF(v_paid_given,N(in_pay_with_return),MAX(v_tc_K,0))")
    add("r24", "Line 24 total paid and adjusted", "=v_r20+v_r21+v_r22+v_r23")
    add("r25", "Line 25 deficit or (excess)", "=v_r19-v_r24")
    add("ack.income", "Acknowledgement: total income", "=v_r11")
    add("ack.paid", "Acknowledgement: total tax paid", "=v_r24")

    for i, key in enumerate(("food_clothing", "housing", "transport", "utility", "education", "travel", "festival",
                             "tax_paid", "loan_interest"), start=1):
        if key == "tax_paid":
            add(f"lb.r{i}", "IT-10BB tax paid", '=IF(in_lf_tax_paid&""<>"",N(in_lf_tax_paid),v_r24)')
        else:
            add(f"lb.r{i}", f"IT-10BB {key}", f"=N(in_lf_{key})")
    for i, key in enumerate(("food_clothing", "housing", "transport", "utility", "education", "travel", "festival",
                             "tax_paid", "loan_interest"), start=1):
        add(f"lb.c{i}", f"IT-10BB note {key}", f'=in_lf_{key}_note&""')
    add("lb.total", "IT-10BB total", "=SUM({lb.r1}:{lb.r9})")
    add("lb.total2", "IT-10BB total (comments column)", "=v_lb_total")
    add("expenses", "Total expenses and losses", "=v_lb_total+v_other_out")

    add("w.1a", "IT-10B 1(a) total income", "=v_r11")
    add("w.1b", "IT-10B 1(b) tax-exempt income", "=v_r26")
    add("w.1c", "IT-10B 1(c) receipt of gift and others",
        '=IF(v_mode="gift",v_non_cash+v_cash_in-v_w_liab+v_expenses-v_prev-v_r11-v_r26,v_gift_in)')
    add("w.gift_label", "Label after 1(c)", '=IF(in_w_gift_label&""<>"",in_w_gift_label&"",IF(v_first_return,"(First Return)",""))')
    add("w.src", "Total sources of fund", "=v_w_1a+v_w_1b+v_w_1c")
    add("w.2", "Net wealth last year", '=IF(v_prev_given,v_prev,"")')
    add("w.3", "Sources plus last year's net wealth", "=v_w_src+v_prev")
    add("w.4a", "Lifestyle expenses", "=v_lb_total")
    add("w.4b", "Gifts, expenses, losses not in IT-10BB", "=v_other_out")
    add("w.exp", "Total expenses and losses", "=v_w_4a+v_w_4b")
    add("w.5", "Net wealth at 30 June", "=v_w_3-v_w_exp")
    add("w.7", "Gross wealth", "=v_w_5+v_w_liab")
    add("a.k", "Cash, bank and funds", '=IF(v_mode="cash",v_w_5+v_w_liab-v_non_cash,v_cash_in)')
    add("a.total", "Total assets", "=v_non_cash+v_a_k")

    add("chk_name", "Check: name", '=IF(v_name="","Enter the taxpayer\'s name.","")')
    add("chk_tin", "Check: TIN", '=IF(LEN(v_tin)<>12,"The TIN must have 12 digits.","")')
    add("chk_gift", "Check: gift", '=IF(AND(v_mode="gift",v_w_1c<0),"Receipt of gift and others is negative ("&TEXT(v_w_1c,"#,##0")'
                                    '&"): the assets are smaller than income + last year\'s wealth - expenses.","")')
    add("chk_cash", "Check: cash", '=IF(AND(v_mode="cash",v_a_k<0),"Cash and bank is negative ("&TEXT(v_a_k,"#,##0")'
                                    '&"): expenses and assets exceed the sources of fund. Add a source or check the figures.","")')
    add("chk_balance", "Check: IT-10B balance", '=IF(ABS(v_a_total-v_w_7)>=1,"IT-10B does not balance: total assets "'
                                                  '&TEXT(v_a_total,"#,##0")&" vs gross wealth "&TEXT(v_w_7,"#,##0")&".","")')
    add("chk_biz", "Check: Schedule 4", '=IF(AND(v_biz_present,ABS(v_s4_r16-v_s4_r10)>=1),"Schedule 4 balance sheet does not balance.","")')
    add("chk_deficit", "Check: deficit", '=IF(v_r25>0,"Deficit of Tk "&TEXT(v_r25,"#,##0")&": payments are less than the tax payable.",'
                                          'IF(v_r25<0,"Excess payment of Tk "&TEXT(-v_r25,"#,##0")&" (refundable).",""))')
    add("chk_min", "Check: minimum tax", '=IF(AND(v_tc_E>v_tc_D,v_tc_F>0),"Minimum tax applies: tax after rebate is "&TEXT(v_tc_D,"#,##0")'
                                          '&" but the minimum is "&TEXT(v_tc_E,"#,##0")&".","")')
    add("chk_rules", "Rules caution", '="Tax rules ("&in_profile&"): "&p_warning')
    return R


def vname(key: str) -> str:
    return "v_" + key.replace(".", "_")


def build_calc(wb):
    ws = wb.create_sheet("Calc")
    print_wide(ws)
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 2), ("B", 22), ("C", 52), ("D", 22)):
        ws.column_dimensions[col].width = w * CHARS
    ws["B1"] = "Calculation"
    ws["B1"].font = Font(name="Times New Roman", size=18, bold=True, color=GREEN)
    ws["B2"] = "Every figure of the return. Column D holds the formula; names v_... point to it."
    ws["B2"].font = Font(name="Arial", size=10)
    rows = calc_rows()
    where = {}
    r = 4
    for key, _, _ in rows:
        where[key] = r
        r += 1
    for key, desc, formula in rows:
        r = where[key]
        # {key} or {key}:{key} placeholders -> cell references (for SUM ranges)
        import re
        f = re.sub(r"\{([^}]+)\}", lambda m: f"D{where[m.group(1)]}", formula)
        ws.cell(r, 2, key).font = Font(name="Arial", size=9, color="5D6B64")
        ws.cell(r, 3, desc).font = Font(name="Arial", size=10)
        c = ws.cell(r, 4, f)
        c.font = Font(name="Arial", size=10)
        c.number_format = MONEY
        c.alignment = Alignment(horizontal="right")
        add_name(wb, vname(key), f"Calc!$D${r}")
    ws.protection.sheet = True
    return ws, {k for k, _, _ in rows}


# =================================================================== print sheets
def page_items(page_no: int, page: dict):
    """Static labels and live fields of one page as layout Items."""
    items = []
    # Labels: words on one baseline in the same style join into runs.
    runs = []
    for w in sorted(page["words"], key=lambda w: (w["base"], w["x"])):
        word = (w["t"], w["x"], w["x"] + text_width(w["t"], w["font"], w["size"]))
        for run in runs:
            if (abs(run["base"] - w["base"]) < 0.4 and run["font"] == w["font"] and run["size"] == w["size"]
                    and run["c"] == w["c"] and 0 <= w["x"] - run["words"][-1][2] <= 0.6 * w["size"]):
                run["words"].append(word)
                break
        else:
            runs.append({"words": [word], "base": w["base"], "font": w["font"], "size": w["size"], "c": w["c"]})

    def split(words, space):
        """A gap much wider than the others on its line parts two labels (a line number and its text)."""
        gaps = [b[1] - a[2] for a, b in zip(words, words[1:])]
        if not gaps:
            return [words]
        i = max(range(len(gaps)), key=gaps.__getitem__)
        rest = sorted(gaps[:i] + gaps[i + 1:])
        if gaps[i] > 1.6 * space and (not rest or gaps[i] > 1.8 * rest[len(rest) // 2]):
            return split(words[:i + 1], space) + split(words[i + 1:], space)
        return [words]

    for run in runs:
        c = run["c"]
        color = (c[0], c[0], c[0]) if len(c) == 1 else tuple(c[:3])
        space = text_width(" ", run["font"], run["size"])
        for words in split(run["words"], space):
            # A line of a justified paragraph has its spaces stretched: spread it the same way.
            gaps = sorted(b[1] - a[2] for a, b in zip(words, words[1:]))
            stretched = len(gaps) >= 3 and gaps[len(gaps) // 2] > 1.12 * space and gaps[0] > 0.9 * space
            items.append(Item(x=words[0][1], base=run["base"], size=run["size"], role=run["font"],
                              align="d" if stretched else "l", width=words[-1][2] - words[0][1],
                              value=" ".join(t for t, _, _ in words), color=color))

    # Runs that touch (a bold link inside a sentence) share one cell as rich text: separate cells
    # would cut the first one off at the second's edge.
    items.sort(key=lambda it: (it.base, it.x))
    labels = []
    for it in items:
        prev = next((o for o in reversed(labels) if abs(o.base - it.base) < 0.5), None)
        if (prev is not None and prev.align == it.align == "l"
                and -0.5 < it.x - (prev.x + prev.width) < 0.8):
            if not prev.rich:
                prev.rich = [(prev.value, prev.role, prev.color)]
            prev.rich.append((it.value, it.role, it.color))
            prev.value += it.value
            prev.width = it.x + it.width - prev.x
        else:
            labels.append(it)
    items = labels

    for f in PAGES[page_no]:
        k, key = f["kind"], f["key"]
        if k == "text":
            formula = f"={vname(key)}"
            if f["lines"] > 1:
                items.append(Item(x=f["x"], base=f["base"], size=f["size"], role=f["font"], align="l", width=f["maxw"],
                                  value=formula, wrap_lines=f["lines"], lead=f["lead"]))
            else:
                items.append(Item(x=f["x"], base=f["base"], size=f["size"], role=f["font"], align=f["align"],
                                  width=f["maxw"], value=formula,
                                  color=tuple(f["color"]) if f.get("color") else (0, 0, 0)))
        elif k == "amount":
            paren = 0.333 * f["size"] if f["font"].startswith(("serif", "sans")) else 0.303 * f["size"]
            formula = f"={vname(key)}" if f["zero"] == "dash" else f'=IF(N({vname(key)})=0,"",{vname(key)})'
            digits = 5.6 if f["font"].startswith("sans") else 5.0  # room for (99,999,999)
            items.append(Item(x=f["right"] + paren, base=f["base"], size=f["size"], role=f["font"], align="r",
                              width=digits * f["size"], value=formula, kind="amount",
                              color=tuple(f["color"]) if f.get("color") else (0, 0, 0),
                              right=f["right"], dash=f.get("dash")))
        elif k == "digits":
            src = vname(key)
            for i, cx in enumerate(f["centers"], start=1):
                items.append(Item(x=cx, base=f["base"], size=f["size"], role=f["font"], align="c", width=f["size"],
                                  value=f"=MID({src},{i},1)"))
        elif k == "check":
            x0, top, x1, bot = f["box"]
            cx = f["at"][0] + 5.28 if f.get("at") else (x0 + x1) / 2
            items.append(Item(x=cx, base=bot - 3.4, size=10.56, role="symbol", align="c", width=11,
                              value=f'=IF({vname(key)},"✔","")', kind="check"))
        elif k == "relation":
            # Two cells, the applicable word in bold. A cell's text cannot run into a filled neighbour,
            # so "/Husband," starts just clear of a bold "Father" (and is all bold for a husband).
            x, b, s = f["x"], f["base"], f["size"]
            father = text_width("Father", "serif-bold", s)
            items.append(Item(x=x, base=b, size=s, role="serif", align="l", width=father, value="Father",
                              cond_bold='v_relation<>"husband"'))
            items.append(Item(x=x + father + 0.6 + PADL, base=b, size=s, role="serif", align="l",
                              width=text_width("/Husband,", "serif-bold", s), value="/Husband,",
                              cond_bold='v_relation="husband"'))
        elif k == "ornaments":
            items.append(Item(x=f["x"], base=f["base"], size=f["size"], role="serif-bold", align="l", width=38,
                              value="=v_a_h_qty"))
            items.append(Item(x=234.77, base=f["base"], size=f["size"], role="serif", align="l", width=28,
                              value="=v_a_h_type_shown"))
    return items


DIGIT_EM = {"serif": 0.5, "calibri": 0.507, "sans": 0.556}
PAREN_EM = {"serif": 0.333, "calibri": 0.303, "sans": 0.333}


def amount_format(it: Item, cell_right: float, s: float) -> str:
    """Accounting format whose padding puts digits and the zero dash where the sample has them."""
    family = it.role.replace("-bold", "")
    size = excel_size(it.size / s) * s
    digit, paren = DIGIT_EM[family] * size, PAREN_EM[family] * size
    end = cell_right - PADR * s
    pad = abs(end - paren - it.right) <= abs(end - it.right)
    dash = it.dash if it.dash is not None else it.right - 2 * digit
    k, zpad = min(((k, z) for k in range(4) for z in (True, False)),
                  key=lambda o: (round(abs(end - (paren if o[1] else 0) - o[0] * digit - dash), 1), o != (2, True)))
    # "_0" (a space as wide as a digit) instead of Excel's usual "??": LibreOffice draws "?" narrower.
    return f'_(* #,##0{"_)" if pad else ""};_(* (#,##0);_(* "-"{"_0" * k}{"_)" if zpad else ""};_(@_)'


def border_style(lw):
    if lw < 0:
        return "dotted" if lw == -2 else "double"
    return "thin" if lw < 1.3 else "medium"


def color_hex(c):
    if c is None:
        return "000000"
    if isinstance(c, (int, float)):
        c = [c]
    if len(c) == 1:
        c = [c[0]] * 3
    return "".join(f"{round(v * 255):02X}" for v in c[:3])


def build_print_sheet(wb, page_no: int):
    page = TEMPLATE["pages"][page_no]
    ws = wb.create_sheet(PAGE_NAMES[page_no])
    items = page_items(page_no, page)
    sc = PAGE_SCALE.get(page_no, 1.0)
    g = build_grid(page, items, TEMPLATE["width"], TEMPLATE["height"], sc)
    xs, ys = g.xs, g.ys
    ncol, nrow = len(xs) - 1, len(ys) - 1
    for c in range(ncol):
        px = round((xs[c + 1] - xs[c]) / sc * 96 / 72)
        ws.column_dimensions[get_column_letter(c + 1)].width = max(px, 1) / DIGIT_PX
    for r in range(nrow):
        ws.row_dimensions[r + 1].height = round((ys[r + 1] - ys[r]) / sc * 96 / 72) * 0.75

    def cell(r, c):
        return ws.cell(r + 1, c + 1)

    # Borders: collect edges per cell, then apply once. A rule along a merged cell goes on the
    # unmerged neighbour across it, so no merged cell carries a border along only part of an edge.
    edges = {}
    merged = {(r, c) for it in g.items if len(it.rows) > 1 or len(it.cols) > 1 for r in it.rows for c in it.cols}

    def edge(r, c, side, style):
        if 0 <= r < nrow and 0 <= c < ncol:
            edges.setdefault((r, c), {})[side] = style

    def bindex(bounds, v):
        return min(range(len(bounds)), key=lambda i: abs(bounds[i] - v))

    for y, x0, x1, lw in g.hsegs:
        ri = bindex(ys, y)
        c0, c1 = bindex(xs, x0), bindex(xs, x1)
        for c in range(c0, c1):
            if ri > 0 and not ((ri - 1, c) in merged and ri < nrow and (ri, c) not in merged):
                edge(ri - 1, c, "bottom", border_style(lw))
            else:
                edge(ri, c, "top", border_style(lw))
    for x, y0, y1, lw in g.vsegs:
        ci = bindex(xs, x)
        r0, r1 = bindex(ys, y0), bindex(ys, y1)
        for r in range(r0, r1):
            if ci > 0 and not ((r, ci - 1) in merged and ci < ncol and (r, ci) not in merged):
                edge(r, ci - 1, "right", border_style(lw))
            else:
                edge(r, ci, "left", border_style(lw))
    for x0, top, x1, bot, fill in g.fills:
        c0, c1, r0, r1 = bindex(xs, x0), bindex(xs, x1), bindex(ys, top), bindex(ys, bot)
        for r in range(r0, r1):
            for c in range(c0, c1):
                cell(r, c).fill = PatternFill("solid", fgColor=color_hex(fill))

    merges = []
    for it in g.items:
        target = cell(it.rows[0], it.cols[0])
        target.value = it.value
        name, bold = EXCEL_FONT[it.role]
        target.font = Font(name=name, size=excel_size(it.size / sc), bold=bold, color=color_hex(it.color),
                           underline="single" if it.underline else None)
        if it.rich:
            target.value = CellRichText([
                TextBlock(InlineFont(rFont=EXCEL_FONT[role][0], sz=excel_size(it.size / sc), b=EXCEL_FONT[role][1],
                                     color=color_hex(color)), text)
                for text, role, color in it.rich])
        if it.wrap_lines > 1:
            target.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        elif it.kind == "amount":
            target.alignment = Alignment(vertical="bottom")
            target.number_format = amount_format(it, xs[it.cols[-1] + 1], sc)
        else:
            # A field gets a cell as wide as its slot on the page; longer text shrinks to fit.
            target.alignment = Alignment(horizontal={"l": "left", "r": "right", "c": "center", "d": "distributed"}[it.align],
                                         vertical="bottom",
                                         shrink_to_fit=it.value.startswith("=") and len(it.cols) > 1)
        if it.cond_bold:
            ws.conditional_formatting.add(target.coordinate, FormulaRule(formula=[it.cond_bold], font=Font(bold=True)))
        if len(it.rows) > 1 or len(it.cols) > 1:
            merges.append((it.rows[0], it.cols[0], it.rows[-1], it.cols[-1]))

    for (r, c), sides in edges.items():
        cur = cell(r, c).border
        cell(r, c).border = Border(
            left=Side(style=sides["left"]) if "left" in sides else cur.left,
            right=Side(style=sides["right"]) if "right" in sides else cur.right,
            top=Side(style=sides["top"]) if "top" in sides else cur.top,
            bottom=Side(style=sides["bottom"]) if "bottom" in sides else cur.bottom)
    for r0, c0, r1, c1 in merges:
        ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1 + 1, end_column=c1 + 1)
        # merge_cells copies the top-left border onto the edges: put back the computed ones
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if (r, c) in edges:
                    sides = edges[(r, c)]
                    cur = cell(r, c).border
                    cell(r, c).border = Border(
                        left=Side(style=sides["left"]) if "left" in sides else cur.left,
                        right=Side(style=sides["right"]) if "right" in sides else cur.right,
                        top=Side(style=sides["top"]) if "top" in sides else cur.top,
                        bottom=Side(style=sides["bottom"]) if "bottom" in sides else cur.bottom)

    for im in g.images:
        img = XLImage(str(EMBLEM))
        c = max(i for i in range(ncol) if xs[i] <= im["x0"] + 0.01)
        r = max(i for i in range(nrow) if ys[i] <= im["top"] + 0.01)
        emu = 12700 / sc
        marker = AnchorMarker(col=c, colOff=int((im["x0"] - xs[c]) * emu), row=r, rowOff=int((im["top"] - ys[r]) * emu))
        size = XDRPositiveSize2D(int((im["x1"] - im["x0"]) * emu), int((im["bot"] - im["top"]) * emu))
        img.anchor = OneCellAnchor(_from=marker, ext=size)
        ws.add_image(img)

    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    if sc == 1.0:
        # Full-size pages fit exactly; fit-to-page only steps in if a printer needs a hair more.
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
    else:
        ws.page_setup.scale = round(sc * 100)
    W, H = TEMPLATE["width"], TEMPLATE["height"]
    # The left and top margins place the page; the right and bottom ones leave a little slack.
    ws.page_margins = PageMargins(left=xs[0] / 72, top=ys[0] / 72, right=max((W - xs[-1]) / 72 - 0.08, 0.05),
                                  bottom=max((H - ys[-1]) / 72 - 0.08, 0.05), header=0, footer=0)
    ws.print_area = f"A1:{get_column_letter(ncol)}{nrow}"
    ws.protection.sheet = True
    return g


def main(out: str, example: str | None = None):
    global EXAMPLE
    if example:
        EXAMPLE = json.loads(Path(example).read_text(encoding="utf-8"))
    wb = Workbook()
    wb._fonts = IndexedList([NORMAL_FONT])
    wb._named_styles["Normal"].font = NORMAL_FONT
    build_input(wb)
    build_rules(wb)
    _, calc_keys = build_calc(wb)
    needed = {f["key"] for page in PAGES for f in page if f["kind"] in ("text", "amount", "digits", "check")}
    missing = sorted(k for k in needed if k not in calc_keys)
    if missing:
        raise SystemExit(f"Calc sheet lacks: {missing}")
    report = []
    for i in range(12):
        g = build_print_sheet(wb, i)
        report.append((PAGE_NAMES[i], len(g.xs) - 1, len(g.ys) - 1, g.problems))
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    for name, nc, nr, probs in report:
        print(f"{name}: {nc} columns x {nr} rows" + (f", {len(probs)} problems: {probs[:3]}" if probs else ""))
    print("saved", out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(HERE / "BD-e-Return-Calculator.xlsx"),
         sys.argv[2] if len(sys.argv) > 2 else None)
